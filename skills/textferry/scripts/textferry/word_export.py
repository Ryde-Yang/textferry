# SPDX-License-Identifier: AGPL-3.0-only
"""Optional Word export with native or explicit image-compatible equations."""
from __future__ import annotations

import json
from copy import deepcopy
from io import BytesIO
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from urllib.parse import unquote, urlsplit
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from docx.image.image import Image as WordImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .math_markdown import normalize_math_delimiters


def pandoc_path():
    try:
        import pypandoc
    except ImportError:
        raise ValueError('Word 导出组件尚未安装，请让 Agent 安装 word 可选依赖后继续。') from None
    binary = Path(pypandoc.__file__).resolve().parent / 'files' / ('pandoc.exe' if os.name == 'nt' else 'pandoc')
    if not binary.is_file():
        raise ValueError('未找到随 Word 组件安装的转换程序，请安装 pypandoc_binary。')
    return binary


def walk(value):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from walk(child)
    elif isinstance(value, list):
        for child in value:
            yield from walk(child)


def _tex_group(text, start, opening, closing):
    while start < len(text) and text[start].isspace():
        start += 1
    if start >= len(text) or text[start] != opening:
        raise ValueError('长等号标注的括号不完整，请核查公式。')
    cursor, depth, braces = start + 1, 1, 0
    while cursor < len(text):
        char = text[cursor]
        if char == '\\':
            cursor += 2
            continue
        if opening == '[':
            braces += (char == '{') - (char == '}')
        if not braces:
            if char == opening:
                depth += 1
            elif char == closing:
                depth -= 1
                if depth == 0:
                    return text[start + 1:cursor], cursor + 1
        cursor += 1
    raise ValueError('长等号标注的括号不完整，请核查公式。')


def portable_tex(text):
    """Keep annotated equalities editable with Pandoc's supported TeX subset."""
    stack, pieces, previous = [], [], 0
    for token in re.finditer(r'\\(begin|end)\{([A-Za-z*]+)\}|(?<!\\)&', text):
        pieces.append(text[previous:token.start()])
        if token[1] == 'begin':
            stack.append(token[2])
            pieces.append(r'\begin{gathered}' if token[2] in {'aligned', 'aligned*'} else token[0])
        elif token[1] == 'end':
            pieces.append(r'\end{gathered}' if token[2] in {'aligned', 'aligned*'} else token[0])
            if stack:
                stack.pop()
        elif not stack or stack[-1] not in {'aligned', 'aligned*'}:
            pieces.append(token[0])
        previous = token.end()
    pieces.append(text[previous:])
    text = ''.join(pieces)
    result, cursor = [], 0
    for match in re.finditer(r'\\xlongequal\b', text):
        if match.start() < cursor:
            continue
        result.append(text[cursor:match.start()])
        position = match.end()
        while position < len(text) and text[position].isspace():
            position += 1
        bottom = None
        if position < len(text) and text[position] == '[':
            bottom, position = _tex_group(text, position, '[', ']')
        top, cursor = _tex_group(text, position, '{', '}')
        equals = r'\overset{' + top + '}{=}'
        result.append(r'\underset{' + bottom + '}{' + equals + '}' if bottom is not None else equals)
    result.append(text[cursor:])
    return ''.join(result)


def reference_document(path, binary):
    # Start with the writer's complete style set. A generic python-docx
    # template leaves Compact/FirstParagraph/FootnoteReference undefined;
    # notably, LibreOffice can detach table cells when Compact is dangling.
    reference = subprocess.run(
        [str(binary), '--print-default-data-file=reference.docx'],
        capture_output=True, check=False,
    )
    if reference.returncode:
        raise ValueError('Word 模板准备失败：' + reference.stderr.decode('utf-8', errors='replace')[:400])
    writer_reference = Document(BytesIO(reference.stdout))
    doc = Document()
    known_styles = {style.style_id for style in doc.styles}
    for style in writer_reference.styles:
        if style.style_id not in known_styles:
            doc.styles.element.append(deepcopy(style.element))
            known_styles.add(style.style_id)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.3)
    for name, size in [('Normal', 11), ('Body Text', 11), ('First Paragraph', 11), ('Compact', 11), ('Title', 18), ('Heading 1', 15), ('Heading 2', 13), ('Heading 3', 12), ('Footnote Text', 9)]:
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, 1)
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            style.element.get_or_add_rPr().append(fonts)
        fonts.set(qn('w:eastAsia'), 'Songti SC' if sys.platform == 'darwin' else ('SimSun' if os.name == 'nt' else 'Noto Serif CJK SC'))
        for attribute in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme'):
            fonts.attrib.pop(qn('w:' + attribute), None)
        style.paragraph_format.line_spacing = 1.35
        style.paragraph_format.space_after = Pt(6)
    for name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        paragraph_style = doc.styles[name].element.get_or_add_pPr()
        for border in list(paragraph_style.findall(qn('w:pBdr'))):
            paragraph_style.remove(border)
        doc.styles[name].font.bold = True
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.comments = ''
    doc.save(path)


def compatible_math_images(nodes, resources):
    """Render a batch once; keep TeX in image alt text and sidecar metadata."""
    selected = [node for node in nodes if node['c'][0]['t'] == 'DisplayMath'
                or re.search(r'\\begin\{[^}]*?(?:array|matrix|cases|aligned|gathered|split)[^}]*\}|\\(?:sum|prod|int|lim)(?![A-Za-z])', node['c'][1])]
    if not selected:
        return {'images': 0, 'scaled': 0}
    node_binary = os.environ.get('TEXTFERRY_NODE') or os.environ.get('NODE') or shutil.which('node')
    if not node_binary:
        raise ValueError('Word 兼容公式模式需要本地 Node.js，以及 mathjax-full、sharp；请让 Agent 先完成公式组件准备。')
    directory = resources / 'word-math'
    if directory.is_symlink() or not directory.resolve().is_relative_to(resources):
        raise ValueError('公式图片缓存必须在任务目录内，不能使用外部链接。')
    directory.mkdir(exist_ok=True)
    for path in directory.iterdir():
        if path.is_symlink():
            raise ValueError('公式图片缓存中不能含有符号链接，请使用任务目录内的普通文件。')
    unique = list(dict.fromkeys((node['c'][1], node['c'][0]['t'] == 'DisplayMath') for node in selected))
    request = {'directory': str(directory), 'equations': [
        {'id': str(index), 'tex': tex, 'display': display} for index, (tex, display) in enumerate(unique)]}
    try:
        result = subprocess.run([node_binary, str(Path(__file__).with_name('render_word_math.cjs'))],
                                input=json.dumps(request), text=True, capture_output=True, check=False, cwd=resources)
    except OSError as error:
        raise ValueError('本地公式渲染程序无法启动，请检查 TEXTFERRY_NODE。') from error
    if result.returncode:
        raise ValueError('Word 兼容公式渲染失败；请检查本地 mathjax-full、sharp 组件或公式，未改用不完整公式：' + result.stderr[:700])
    try:
        records = json.loads(result.stdout)['equations']
        if len(records) != len(unique):
            raise ValueError('公式图片数量不匹配。')
        by_formula = dict(zip(unique, records))
        scaled = 0
        for node in selected:
            tex, display = node['c'][1], node['c'][0]['t'] == 'DisplayMath'
            record = by_formula[(tex, display)]
            path = Path(record['path']).resolve()
            if not path.is_relative_to(directory.resolve()) or not path.is_file():
                raise ValueError('公式图片必须保存在任务目录内。')
            width, height = float(record['width_cm']), float(record['height_cm'])
            if not (0 < width <= 16.4 and 0 < height <= 18):
                raise ValueError('公式图片尺寸超出页面范围。')
            scaled += record['scale'] < 1
            node.clear()
            node.update({'t': 'Image', 'c': [['', ['word-equation'], [
                ['width', f'{width:.6f}cm'], ['height', f'{height:.6f}cm']]],
                [{'t': 'Str', 'c': tex}], [str(path), '']]})
    except (KeyError, TypeError, json.JSONDecodeError) as error:
        raise ValueError('公式渲染结果不完整，未生成 Word。') from error
    return {'images': len(selected), 'scaled': scaled}


def write_word(text, output, resource_dir, *, equations=None):
    from .progress import emit
    emit("word_export", "正在生成 Word，复用已完成的文字和图片")
    mode = equations or os.environ.get('TEXTFERRY_WORD_EQUATIONS', 'native').strip().lower()
    if mode not in {'native', 'compatible'}:
        raise ValueError('TEXTFERRY_WORD_EQUATIONS 必须是 native 或 compatible。')
    binary = pandoc_path()
    text = normalize_math_delimiters(text)
    # Source numbering must neither disappear for empty items nor restart in
    # Word's automatic lists. Escape labels outside code, retaining their text.
    labelled, fence = [], None
    for line in text.splitlines(keepends=True):
        mark = re.match(r'^ {0,3}(`{3,}|~{3,})', line)
        if mark:
            if fence is None:
                fence = mark[1]
            elif mark[1][0] == fence[0] and len(mark[1]) >= len(fence):
                fence = None
        elif fence is None:
            line = re.sub(r'^( {0,3})(\(\d+\)|\d+[.)])(?=\s|$)',
                          lambda m: m[1] + m[2][:-1] + '\\' + m[2][-1], line)
        labelled.append(line)
    text = ''.join(labelled)
    output, resources = Path(output), Path(resource_dir).resolve()
    parsed = subprocess.run([str(binary), '--from=markdown+tex_math_dollars+tex_math_single_backslash+footnotes+raw_html-markdown_in_html_blocks-raw_tex', '--to=json'], input=text, text=True, capture_output=True, check=False, cwd=resources)
    if parsed.returncode:
        raise ValueError('Word 内容解析失败：' + parsed.stderr[:400])
    tree = json.loads(parsed.stdout)
    # Parse table HTML as document structure; never embed or execute raw HTML.
    for node in list(walk(tree)):
        if node.get('t') == 'RawBlock' and node['c'][0] == 'html' and re.search(r'<table\b', node['c'][1], re.I):
            converted = subprocess.run([str(binary), '--from=html', '--to=json'], input=node['c'][1], text=True, capture_output=True, check=False, cwd=resources)
            if converted.returncode:
                raise ValueError('HTML 表格解析失败，请核查表格结构。')
            blocks = json.loads(converted.stdout)['blocks']
            if not any(item.get('t') == 'Table' for item in walk(blocks)):
                raise ValueError('未能识别 HTML 表格，不能按已转换交付。')
            node.clear()
            node.update({'t': 'Div', 'c': [['', [], []], blocks]})
        elif node.get('t') in {'RawBlock', 'RawInline'} and node['c'][0] == 'html':
            value = node['c'][1]
            kind = node['t']
            node.clear()
            node.update({'t': 'Plain', 'c': [{'t': 'Str', 'c': value}]} if kind == 'RawBlock' else {'t': 'Str', 'c': value})
    if tree.get('blocks') and tree['blocks'][0].get('t') == 'Header' and tree['blocks'][0]['c'][0] == 1:
        header = tree['blocks'][0]['c'][2]
        tree['blocks'][0] = {'t': 'Div', 'c': [['', [], [['custom-style', 'Title']]], [{'t': 'Para', 'c': header}]]}
    math_nodes = [node for node in walk(tree) if node.get('t') == 'Math']
    for node in walk(tree):
        if node.get('t') == 'Image':
            link = node['c'][-1][0]
            parts = urlsplit(link)
            path = (resources / unquote(parts.path)).resolve()
            if parts.scheme or not path.is_relative_to(resources) or not path.is_file():
                raise ValueError('Word 中的图片必须是任务目录内已有的本地图片，请先核实图片来源。')
            node['c'][-1][0] = str(path)
            attributes = node['c'][0][2]
            if not any(key in {'width', 'height'} for key, _ in attributes) and path.suffix.lower() in {'.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tif', '.tiff'}:
                image = WordImage.from_file(str(path))
                ratio = image.px_width / image.px_height
                width = min(image.width.cm, 16.4, 18 * ratio)
                attributes.extend([['width', f'{width:.4f}cm'], ['height', f'{width / ratio:.4f}cm']])
    compatibility = compatible_math_images(math_nodes, resources) if mode == 'compatible' else {'images': 0, 'scaled': 0}
    native = len(math_nodes) - compatibility['images']
    for node in math_nodes:
        if node.get('t') == 'Math':
            node['c'][1] = portable_tex(node['c'][1])
    with tempfile.TemporaryDirectory(prefix='.word-', dir=resources) as temp:
        temp = Path(temp)
        reference = temp / 'reference.docx'
        reference_document(reference, binary)
        built = temp / 'output.docx'
        result = subprocess.run([str(binary), '--from=json', '--to=docx', '--reference-doc=' + str(reference), '--output=' + str(built)], input=json.dumps(tree, ensure_ascii=False), text=True, capture_output=True, check=False, cwd=resources)
        if result.returncode or 'Could not convert TeX math' in result.stderr:
            raise ValueError('Word 排版未完成，请检查公式或文档结构：' + result.stderr[:500])
        document_style = Document(built)
        paragraphs = document_style.paragraphs
        for index, paragraph in enumerate(paragraphs):
            if index and re.fullmatch(r'\(\d+(?:\.\d+)+\)', paragraph.text.strip()):
                paragraphs[index - 1].paragraph_format.keep_with_next = True
        math_ns = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
        for matrix in document_style.element.iter(math_ns + 'm'):
            rows = matrix.findall(math_ns + 'mr')
            while len(rows) > 1 and not ''.join(rows[-1].itertext()).strip() and all(not len(cell) for cell in rows[-1]):
                matrix.remove(rows.pop())
            columns = max((len(row.findall(math_ns + 'e')) for row in rows), default=0)
            for row in rows:
                while len(row.findall(math_ns + 'e')) < columns:
                    row.append(OxmlElement('m:e'))
                for cell in row.findall(math_ns + 'e'):
                    if not len(cell):
                        run, value = OxmlElement('m:r'), OxmlElement('m:t')
                        value.text = '\u00a0'
                        value.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                        run.append(value)
                        cell.append(run)
        for table in document_style.tables:
            table.autofit = False
            available = int(document_style.sections[0].page_width - document_style.sections[0].left_margin - document_style.sections[0].right_margin)
            width = available // max(1, len(table.columns))
            table_width = table._tbl.tblPr.find(qn('w:tblW'))
            if table_width is None:
                table_width = OxmlElement('w:tblW')
                table._tbl.tblPr.append(table_width)
            table_width.set(qn('w:type'), 'dxa')
            table_width.set(qn('w:w'), str(round(available / 635)))
            if 'Table Grid' in document_style.styles:
                table.style = 'Table Grid'
            for column in table.columns:
                column.width = width
            for row in table.rows:
                for cell in row.cells:
                    cell.width = width
        document_style.core_properties.author = ''
        document_style.core_properties.last_modified_by = ''
        document_style.save(built)
        with ZipFile(built) as archive:
            document = ET.fromstring(archive.read('word/document.xml'))
            actual = len(document.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))
            if actual < native:
                raise ValueError('部分公式未能转换成可编辑的 Word 公式，请先核查这些公式。')
        built.replace(output)
    return {'equations': len(math_nodes), 'native_equations': actual,
            'image_equations': compatibility['images'], 'scaled_equations': compatibility['scaled'],
            'equation_mode': mode, 'file': output.name}
