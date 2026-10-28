# SPDX-License-Identifier: AGPL-3.0-only
"""分页/连续两类人工可读 docx、连续版跨页断句合并，以及分页版/连续版真实脚注写回报告。"""

from __future__ import annotations

import re
from dataclasses import dataclass
from html import unescape
from pathlib import Path
from typing import Iterable, List

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .docx_export import ExportPage, _split_page_text_blocks, _write_failed_page, src_key
from .export_cleanup import clean_reader_text, detect_repeated_noise_lines, iter_blocks_preserving_tables
from .footnote_detection import extract_page_footnotes, materialize_footnote_text
from .ooxml_footnotes import OoxmlFootnote, inject_footnotes


HTML_ROW_RE = re.compile(r"<tr>(.*?)</tr>", re.IGNORECASE | re.DOTALL)
HTML_CELL_RE = re.compile(r"<t[dh]>(.*?)</t[dh]>", re.IGNORECASE | re.DOTALL)
HEADING_RE = re.compile(r"^(第[0-9一二三四五六七八九十百千]+[章节篇卷回].*|[一二三四五六七八九十]+[、．.].*)$")
FLOW_SENTENCE_END_RE = re.compile(r"[。！？!?；;：:）】》」』”’]$")
FLOW_NO_MERGE_START_RE = re.compile(r"^(第[0-9一二三四五六七八九十百千]+[章节篇卷回]|[一二三四五六七八九十]+[、．.]|[0-9]+[、．.])")
FLOW_SHORT_TITLE_RE = re.compile(r"^(总序|目录|前言|后记|序|跋|.*出版社|.*[著编])$")
FOOTNOTE_PLACEHOLDER_RE = re.compile(r"\ue001FN(\d+)\ue002")
_CJK_RANGE_RE = re.compile(r"[\u4e00-\u9fff\u3400-\u4dbf]")


@dataclass(frozen=True)
class ReaderFootnotePageDecision:
    page_no: int
    attempted: bool
    materialized_notes: int
    agent_resolved_notes: int = 0
    agent_rejected_notes: int = 0
    marker_recovered_notes: int = 0
    layout_assisted_notes: int = 0
    reason: str | None = None


@dataclass(frozen=True)
class PreparedReaderFootnote:
    id: int
    placeholder: str
    text: str


@dataclass(frozen=True)
class PreparedReaderPage:
    page_no: int
    text: str
    footnotes: List[PreparedReaderFootnote]
    decision: ReaderFootnotePageDecision


@dataclass(frozen=True)
class ReaderExportReport:
    attempted_pages: int
    materialized_notes: int
    skipped_pages: int
    skipped_reasons: dict[str, int]
    marker_recovered_notes: int
    layout_assisted_notes: int
    agent_resolved_notes: int
    agent_rejected_notes: int
    agent_resolved_pages: int
    agent_rejected_pages: int
    page_decisions: List[ReaderFootnotePageDecision]


def write_reader_docx(
    output_path: Path,
    pages: Iterable[ExportPage],
    *,
    preserve_page_breaks: bool = True,
    include_page_headers: bool = True,
    merge_across_pages: bool = False,
    convert_footnotes: bool = False,
    prepared_pages: dict[int, PreparedReaderPage] | None = None,
    footnote_restart_mode: str = "eachPage",
) -> ReaderExportReport:
    page_list: List[ExportPage] = list(pages)
    repeated_noise_lines = detect_repeated_noise_lines([page.text for page in page_list if page.text.strip()])
    doc = Document()
    footnotes: List[OoxmlFootnote] = []
    footnote_map: dict[str, int] = {}
    page_decisions: List[ReaderFootnotePageDecision] = []
    last_paragraph = None
    last_paragraph_text = ""

    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "宋体"
    normal_style.font.size = Pt(11.5)
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for idx, page in enumerate(page_list):
        if include_page_headers:
            _write_page_header(doc, page)
            last_paragraph = None
            last_paragraph_text = ""

        if page.failed:
            _write_failed_page(doc, page)
            last_paragraph = None
            last_paragraph_text = ""
        else:
            prepared_page = (prepared_pages or {}).get(page.pdf_page_no)
            reader_text = clean_reader_text(page.text, repeated_noise_lines)
            if convert_footnotes and prepared_page is not None:
                reader_text = prepared_page.text
                for footnote in prepared_page.footnotes:
                    footnotes.append(OoxmlFootnote(id=footnote.id, key=str(footnote.id), text=footnote.text))
                    footnote_map[footnote.placeholder] = footnote.id
                page_decisions.append(prepared_page.decision)
            elif convert_footnotes:
                detected = extract_page_footnotes(reader_text, next_id=len(footnotes) + 2)
                if detected.attempted:
                    if detected.resolved_footnotes:
                        reader_text = materialize_footnote_text(
                            detected.body_text,
                            detected.resolved_footnotes,
                            unresolved_lines=detected.unresolved_lines,
                        )
                        for footnote in detected.resolved_footnotes:
                            footnotes.append(
                                OoxmlFootnote(id=footnote.id, key=str(footnote.id), text=footnote.text)
                            )
                            footnote_map[footnote.placeholder] = footnote.id
                        page_decisions.append(
                            ReaderFootnotePageDecision(
                                page_no=page.pdf_page_no,
                                attempted=True,
                                materialized_notes=len(detected.resolved_footnotes),
                                marker_recovered_notes=detected.marker_recovered_notes,
                                layout_assisted_notes=detected.layout_assisted_notes,
                            )
                        )
                    else:
                        page_decisions.append(
                            ReaderFootnotePageDecision(
                                page_no=page.pdf_page_no,
                                attempted=True,
                                materialized_notes=0,
                                reason=detected.skipped_reason or "unknown",
                            )
                        )
            for block_kind, block in iter_blocks_preserving_tables(reader_text):
                if block_kind == "table":
                    _write_html_table(doc, block)
                    last_paragraph = None
                    last_paragraph_text = ""
                    continue
                for paragraph_text in _split_page_text_blocks(block):
                    if (
                        merge_across_pages
                        and last_paragraph is not None
                        and _should_merge_flow_paragraph(last_paragraph_text, paragraph_text)
                    ):
                        last_paragraph_text = f"{last_paragraph_text}{paragraph_text}"
                        _replace_reader_paragraph_with_footnotes(last_paragraph, last_paragraph_text, footnote_map)
                        continue
                    last_paragraph = _write_reader_paragraph(doc, paragraph_text, footnote_map)
                    last_paragraph_text = paragraph_text

        if preserve_page_breaks and idx < len(page_list) - 1:
            doc.add_page_break()
            last_paragraph = None
            last_paragraph_text = ""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    if convert_footnotes and footnotes:
        inject_footnotes(str(output_path), str(output_path), footnotes, restart_mode=footnote_restart_mode)

    attempted_pages = sum(1 for decision in page_decisions if decision.attempted)
    materialized_notes = sum(decision.materialized_notes for decision in page_decisions)
    skipped_reasons: dict[str, int] = {}
    for decision in page_decisions:
        if decision.materialized_notes > 0 or not decision.reason:
            continue
        skipped_reasons[decision.reason] = skipped_reasons.get(decision.reason, 0) + 1
    skipped_pages = sum(1 for decision in page_decisions if decision.materialized_notes == 0 and bool(decision.reason))
    marker_recovered_notes = sum(decision.marker_recovered_notes for decision in page_decisions)
    layout_assisted_notes = sum(decision.layout_assisted_notes for decision in page_decisions)
    agent_resolved_notes = sum(decision.agent_resolved_notes for decision in page_decisions)
    agent_rejected_notes = sum(decision.agent_rejected_notes for decision in page_decisions)
    agent_resolved_pages = sum(1 for decision in page_decisions if decision.agent_resolved_notes > 0)
    agent_rejected_pages = sum(1 for decision in page_decisions if decision.agent_rejected_notes > 0)
    return ReaderExportReport(
        attempted_pages=attempted_pages,
        materialized_notes=materialized_notes,
        skipped_pages=skipped_pages,
        skipped_reasons=skipped_reasons,
        marker_recovered_notes=marker_recovered_notes,
        layout_assisted_notes=layout_assisted_notes,
        agent_resolved_notes=agent_resolved_notes,
        agent_rejected_notes=agent_rejected_notes,
        agent_resolved_pages=agent_resolved_pages,
        agent_rejected_pages=agent_rejected_pages,
        page_decisions=page_decisions,
    )


def _write_page_header(doc: Document, page: ExportPage) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(f"{page.resolved_label} | SRC:{src_key(page.pdf_page_no)}")
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(110, 110, 110)


def _write_reader_paragraph(doc: Document, block: str, footnote_map: dict[str, int]):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    written_runs = _write_runs_with_footnotes(paragraph, block, footnote_map)
    if HEADING_RE.match(block) and len(block) <= 28:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        for run in written_runs:
            run.bold = True
            run.font.size = Pt(13)
    return paragraph


def _write_runs_with_footnotes(paragraph, text: str, footnote_map: dict[str, int]):
    runs = []
    cursor = 0
    for match in FOOTNOTE_PLACEHOLDER_RE.finditer(text):
        placeholder = match.group(0)
        footnote_id = footnote_map.get(placeholder)
        if footnote_id is None:
            if match.start() > cursor:
                runs.append(_add_reader_run(paragraph, text[cursor : match.start()]))
            cursor = match.end()
            continue
        if match.start() > cursor:
            runs.append(_add_reader_run(paragraph, text[cursor : match.start()]))
        _add_footnote_reference(paragraph, footnote_id)
        cursor = match.end()
    if cursor < len(text):
        runs.append(_add_reader_run(paragraph, text[cursor:]))
    if not runs and not footnote_map:
        runs.append(_add_reader_run(paragraph, text))
    return runs


def _add_reader_run(paragraph, text: str):
    run = paragraph.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(11.5)
    return run


def _replace_reader_paragraph_with_footnotes(paragraph, text: str, footnote_map: dict[str, int]) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    for ref in paragraph._element.findall(f".//{qn('w:footnoteReference')}"):
        r_elem = ref.getparent()
        if r_elem is not None:
            r_elem.getparent().remove(r_elem)
    _write_runs_with_footnotes(paragraph, text, footnote_map)


def _should_merge_flow_paragraph(previous: str, current: str) -> bool:
    previous = previous.strip()
    current = current.strip()
    if not previous or not current:
        return False
    if len(previous) < 4 or len(current) < 2:
        return False
    if FLOW_SENTENCE_END_RE.search(previous):
        return False
    if HEADING_RE.match(previous) or HEADING_RE.match(current):
        return False
    if _looks_like_flow_short_title(previous) or _looks_like_flow_short_title(current):
        return False
    if FLOW_NO_MERGE_START_RE.match(current):
        return False
    if _is_latin_only(previous) and _starts_with_cjk(current):
        return False
    return True


def _is_latin_only(text: str) -> bool:
    return bool(text) and not _CJK_RANGE_RE.search(text)


def _starts_with_cjk(text: str) -> bool:
    stripped = text.lstrip()
    return bool(stripped) and bool(_CJK_RANGE_RE.match(stripped[0]))


def _looks_like_flow_short_title(text: str) -> bool:
    text = text.strip()
    if len(text) > 14:
        return False
    if FLOW_SENTENCE_END_RE.search(text):
        return False
    return bool(FLOW_SHORT_TITLE_RE.match(text))


def _add_footnote_reference(paragraph, footnote_id: int) -> None:
    run = paragraph.add_run()
    r_pr = OxmlElement("w:rPr")
    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), "FootnoteReference")
    r_pr.append(r_style)
    valign = OxmlElement("w:vertAlign")
    valign.set(qn("w:val"), "superscript")
    r_pr.append(valign)
    run._r.append(r_pr)
    reference = OxmlElement("w:footnoteReference")
    reference.set(qn("w:id"), str(footnote_id))
    run._r.append(reference)


def _write_html_table(doc: Document, html_table: str) -> None:
    rows = HTML_ROW_RE.findall(html_table)
    parsed_rows = []
    for row in rows:
        cells = [unescape(_strip_html_tags(cell)).strip() for cell in HTML_CELL_RE.findall(row)]
        if cells:
            parsed_rows.append(cells)
    if not parsed_rows:
        _write_reader_paragraph(doc, _strip_html_tags(html_table), {})
        return
    col_count = max(len(row) for row in parsed_rows)
    table = doc.add_table(rows=len(parsed_rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_idx, row in enumerate(parsed_rows):
        for col_idx in range(col_count):
            value = row[col_idx] if col_idx < len(row) else ""
            cell = table.cell(row_idx, col_idx)
            cell.text = FOOTNOTE_PLACEHOLDER_RE.sub("", value)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "宋体"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                    run.font.size = Pt(10.5)
    doc.add_paragraph("")


def _strip_html_tags(text: str) -> str:
    return re.sub(r"<[^>]+>", "", text)
