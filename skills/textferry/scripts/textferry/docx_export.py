# SPDX-License-Identifier: AGPL-3.0-only
"""带 SRC 标记、失败页占位与可人工复核图片的 textferry docx。"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SENTENCE_END_RE = re.compile(r"[。！？!?；;：:）】》」』”’]$")
TOC_LINE_RE = re.compile(r"[.…·]{2,}\s*[0-9ivxlcdmIVXLCDM一二三四五六七八九十百千]+\s*$")
HEADING_PREFIX_RE = re.compile(r"^(第[0-9一二三四五六七八九十百千]+[章节篇卷回]|[（(]?[0-9一二三四五六七八九十]+[)）、.．])")


@dataclass(frozen=True)
class ExportPage:
    pdf_page_no: int
    resolved_label: str
    text: str
    failed: bool
    image_path: Optional[Path] = None
    failure_reason: Optional[str] = None


def src_key(pdf_page_no: int) -> str:
    return f"pdf_{pdf_page_no:04d}"


_CJK_CHAR_RE = re.compile(r"[\u4e00-\u9fff\u3400-\u4dbf]")


def _should_merge_line(previous: str, current: str) -> bool:
    if not previous or not current:
        return False
    if SENTENCE_END_RE.search(previous):
        return False
    if TOC_LINE_RE.search(previous):
        return False
    if HEADING_PREFIX_RE.match(current):
        return False
    if len(previous.strip()) < 4 or len(current.strip()) < 4:
        return False
    if not _CJK_CHAR_RE.search(previous) and _CJK_CHAR_RE.match(current.lstrip()):
        return False
    return True


def _split_page_text_blocks(text: str) -> List[str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    raw_blocks = [block for block in re.split(r"\n\s*\n+", normalized) if block.strip()]
    blocks: List[str] = []

    for raw_block in raw_blocks:
        lines = [line.strip() for line in raw_block.splitlines() if line.strip()]
        if not lines:
            continue

        current = lines[0]
        for line in lines[1:]:
            if _should_merge_line(current, line):
                current = f"{current}{line}"
            else:
                blocks.append(current.strip())
                current = line
        blocks.append(current.strip())

    return [block for block in blocks if block]


def _write_failed_page(doc: Document, page: ExportPage) -> None:
    placeholder = doc.add_paragraph("【该页识别失败，已保留占位；可在任务详情中补跑失败页】")
    placeholder.runs[0].italic = True
    if page.failure_reason:
        reason = doc.add_paragraph(f"失败原因：{page.failure_reason}")
        reason.runs[0].italic = True
    if page.image_path and page.image_path.exists():
        image_paragraph = doc.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.add_run().add_picture(str(page.image_path), width=Inches(5.8))
