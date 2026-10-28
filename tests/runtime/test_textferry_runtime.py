"""主流程、结构化诊断、失败补跑、正文确认、恢复、机器版结构化页块、分页/连续产物、flow 页码噪声清理、按页脚注编号、agent 冲突页决策、保守真实脚注策略、provider响应解析容错、去噪模型配置与百度降并发的回归测试。"""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import httpx
from docx import Document

from textferry.config import _float_env, get_model_registry, get_textferry_settings
from textferry.normalize import normalize_ocr_text
from textferry.ocr_client import BaiduAsyncProvider, OcrClientError, SiliconFlowProvider
from textferry.service import TextFerryService
from textferry.types import JobStatus, PageMode, PageStatus, ScriptTarget


def build_service() -> TextFerryService:
    get_textferry_settings.cache_clear()
    get_model_registry.cache_clear()
    return TextFerryService()


def test_float_env_parses_and_falls_back(monkeypatch) -> None:
    monkeypatch.setenv("TEXTFERRY_FLOAT_OK", "12.5")
    monkeypatch.setenv("TEXTFERRY_FLOAT_BAD", "oops")

    assert _float_env("TEXTFERRY_FLOAT_OK", 1.0) == 12.5
    assert _float_env("TEXTFERRY_FLOAT_BAD", 1.0) == 1.0
    assert _float_env("TEXTFERRY_FLOAT_MISSING", 1.0) == 1.0


def test_settings_read_does_not_create_data_dir(monkeypatch, tmp_path: Path) -> None:
    data_dir = tmp_path / "config-only-data"
    monkeypatch.setenv("TEXTFERRY_DATA_DIR", str(data_dir))
    get_textferry_settings.cache_clear()

    settings = get_textferry_settings()

    assert settings.data_dir == data_dir.resolve()
    assert not data_dir.exists()


def test_default_data_dir_uses_cache_dir(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setattr(Path, "home", lambda: tmp_path)
    monkeypatch.delenv("TEXTFERRY_DATA_DIR", raising=False)
    monkeypatch.delenv("XDG_CACHE_HOME", raising=False)
    get_textferry_settings.cache_clear()

    settings = get_textferry_settings()

    assert settings.data_dir == (Path.home() / ".cache" / "textferry").resolve()
    assert not settings.data_dir.exists()


def test_default_data_dir_honors_user_cache_location(monkeypatch, tmp_path: Path) -> None:
    cache = tmp_path / "user-cache"
    monkeypatch.setenv("XDG_CACHE_HOME", str(cache))
    monkeypatch.delenv("TEXTFERRY_DATA_DIR", raising=False)
    get_textferry_settings.cache_clear()

    settings = get_textferry_settings()

    assert settings.data_dir == (cache / "textferry").resolve()
    assert not settings.data_dir.exists()


def test_normalize_strips_classical_book_running_headers() -> None:
    raw_text = (
        "孟子注疏\n"
        "前文说明。\n"
        "正文甲。\n"
        "孟子注疏解經卷第六下\n"
        "正文甲。\n"
        "卷第七上\n"
        "正文乙。\n"
        "三三一\n"
    )

    normalized = normalize_ocr_text(raw_text, ScriptTarget.SIMP)

    assert "孟子注疏\n" not in normalized
    assert "孟子注疏解经卷第六下" not in normalized
    assert "卷第七上" not in normalized
    assert "前文说明。" in normalized
    assert "正文甲。" in normalized
    assert "正文乙。" in normalized
    assert "三三一" in normalized


def test_normalize_collapses_pathological_repeated_cjk_clauses() -> None:
    repeated_clause = "足在下，人之所賤，水濁而濯纓，則清者人之所貴也；"
    raw_text = (
        "滄浪之水濁兮，可以濯我足者。\n"
        + repeated_clause * 12
        + "\n而濯足，則濁者人之所賤也。"
    )

    normalized = normalize_ocr_text(raw_text, ScriptTarget.SIMP)
    simplified_clause = "足在下，人之所贱，水浊而濯缨，则清者人之所贵也；"

    assert normalized.count(simplified_clause) == 1
    assert "而濯足，则浊者人之所贱也。" in normalized


def test_create_job_and_run_to_final_ready(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "sample.pdf", 3)
    service = build_service()

    created = service.create_job_from_pdf_path(
        pdf_path=pdf_path,
        page_mode=PageMode.PDF,
        model_option="paddleocr_vl",
        script_target=ScriptTarget.SIMP,
    )
    assert created.status == JobStatus.CREATED

    result = service.run_job(created.job_id, provider="baidu_async")
    assert result["status"] == JobStatus.FINAL_READY.value
    assert result["progress"]["failedPagesCount"] == 0

    job_dir = service.store.job_dir(created.job_id)
    assert (job_dir / "job.json").exists()
    assert (job_dir / "outputs" / "reader.docx").exists()
    assert (job_dir / "outputs" / "machine.md").exists()

    page = service.store.get_page(created.job_id, 1)
    assert page["resolvedLabel"]["mode"] == "pdf"
    assert page["resolvedLabel"]["value"] == 1
    assert page["resolvedLabel"]["fallback"] is False
    assert page["status"] == PageStatus.DONE.value

    events = service.store.read_events(created.job_id)
    assert events
    assert set(events[0].keys()) == {"timestamp", "event", "jobId", "page", "detail"}

    diagnostics = service.store.read_diagnostics(created.job_id)
    assert isinstance(diagnostics, list)


def test_docx_disabled_artifacts_fallback_to_machine_md(textferry_env: Path, make_pdf, monkeypatch) -> None:
    monkeypatch.setenv("TEXTFERRY_ENABLE_READER_DOCX", "0")
    get_textferry_settings.cache_clear()
    get_model_registry.cache_clear()

    pdf_path = make_pdf(textferry_env / "docx-off.pdf", 3)
    service = TextFerryService()
    created = service.create_job_from_pdf_path(
        pdf_path=pdf_path,
        page_mode=PageMode.PDF,
        model_option="paddleocr_vl",
        script_target=ScriptTarget.SIMP,
    )
    result = service.run_job(created.job_id, provider="baidu_async")
    assert result["status"] == JobStatus.FINAL_READY.value

    job_dir = service.store.job_dir(created.job_id)
    outputs_dir = job_dir / "outputs"
    assert (outputs_dir / "machine.md").exists()
    assert not (outputs_dir / "reader.docx").exists()
    assert not (outputs_dir / "reader-flow.docx").exists()
    assert not (outputs_dir / "reader-paged.docx").exists()

    job = service.store.get_job(created.job_id)
    artifacts = job["artifacts"]
    assert artifacts["best"] == "outputs/machine.md"
    assert artifacts["final"] == "outputs/machine.md"
    assert artifacts["readerDocx"] is None
    assert artifacts["readerFlowDocx"] is None
    assert artifacts["readerPagedDocx"] is None


def test_retry_failed_only_requeues_failed_pages(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "retry.pdf", 3)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    original_mark_done = service._mark_page_done_from_raw

    def flaky_mark_done(**kwargs):
        if kwargs["page_no"] == 2:
            raise OcrClientError("empty OCR response")
        return original_mark_done(**kwargs)

    service._mark_page_done_from_raw = flaky_mark_done  # type: ignore[method-assign]
    first = service.run_job(created.job_id, provider="baidu_async")
    assert first["status"] == JobStatus.PARTIAL_READY.value
    assert [item["pdfPageNo"] for item in first["failedPages"]] == [2]

    retry_result = service.retry_failed(created.job_id)
    assert retry_result["retryQueuedPages"] == 1

    events_before = len(service.store.read_events(created.job_id))
    service._mark_page_done_from_raw = original_mark_done  # type: ignore[method-assign]
    second = service.run_job(created.job_id, provider="baidu_async")
    assert second["status"] == JobStatus.FINAL_READY.value

    page1 = service.store.get_page(created.job_id, 1)
    page2 = service.store.get_page(created.job_id, 2)
    page3 = service.store.get_page(created.job_id, 3)
    assert page1["status"] == PageStatus.DONE.value
    assert page2["status"] == PageStatus.DONE.value
    assert page3["status"] == PageStatus.DONE.value

    new_events = service.store.read_events(created.job_id)[events_before:]
    processed_pages = [event["page"] for event in new_events if event["event"] == "page_processing"]
    assert processed_pages == [2]


def test_retry_failed_without_failed_pages_returns_current_status(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "retry-none.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)
    final = service.run_job(created.job_id, provider="baidu_async")
    assert final["status"] == JobStatus.FINAL_READY.value

    retry_result = service.retry_failed(created.job_id)

    assert retry_result["status"] == JobStatus.FINAL_READY.value
    assert retry_result["retryQueuedPages"] == 0
    assert service.get_job_view(created.job_id)["status"] == JobStatus.FINAL_READY.value


def test_body_mode_waits_for_confirmation_then_finishes(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "body.pdf", 3)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.BODY)

    waiting = service.run_job(created.job_id, provider="baidu_async")
    assert waiting["status"] == JobStatus.WAITING_BODY_START.value
    assert waiting["artifacts"]["partial"]
    assert waiting["artifacts"]["best"]

    confirmed = service.confirm_body_start(created.job_id, 2)
    assert confirmed["bodyStartPdfPage"] == 2

    final = service.run_job(created.job_id, provider="baidu_async")
    assert final["status"] == JobStatus.FINAL_READY.value
    page2 = service.store.get_page(created.job_id, 2)
    assert page2["resolvedLabel"]["mode"] == "body"
    assert page2["resolvedLabel"]["value"] == 1


def test_confirm_body_start_rejects_non_body_job_without_changing_status(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "confirm-non-body.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)
    final = service.run_job(created.job_id, provider="baidu_async")
    assert final["status"] == JobStatus.FINAL_READY.value

    try:
        service.confirm_body_start(created.job_id, 1)
    except ValueError as exc:
        assert "body" in str(exc).lower()
    else:
        raise AssertionError("confirm_body_start should reject non-body jobs")

    assert service.get_job_view(created.job_id)["status"] == JobStatus.FINAL_READY.value


def test_recover_processing_pages_on_resume(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "resume.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    service.store.mark_page_processing(created.job_id, 1)
    assert service.store.get_page(created.job_id, 1)["status"] == PageStatus.PROCESSING.value

    result = service.run_job(created.job_id, provider="baidu_async")
    assert result["status"] == JobStatus.FINAL_READY.value
    assert service.store.get_page(created.job_id, 1)["status"] == PageStatus.DONE.value


def test_quality_gate_records_diagnostics(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "quality.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(
        pdf_path=pdf_path,
        page_mode=PageMode.PDF,
        model_option="paddleocr_vl",
    )

    async def low_quality(*, expected_pages: int, **kwargs):
        noise = "<|begin>" * 25 + "<nl>" * 10
        if expected_pages == 1:
            return {1: noise}
        return {1: noise, 2: "这是正常的文本\n第 2 页"}

    service._baidu.ocr_pdf_pages_baidu_async = low_quality  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")
    assert result["status"] == JobStatus.PARTIAL_READY.value
    assert result["failedPages"][0]["pdfPageNo"] == 1

    diagnostics = service.store.read_diagnostics(created.job_id)
    assert diagnostics
    assert diagnostics[0]["type"] == "quality_gate"
    assert set(diagnostics[0].keys()) == {"timestamp", "type", "jobId", "message", "context"}


def test_baidu_timeout_split_records_diagnostics(textferry_env: Path, make_pdf, monkeypatch) -> None:
    monkeypatch.setenv("TEXTFERRY_BAIDU_BATCH_SIZE", "10")
    pdf_path = make_pdf(textferry_env / "split.pdf", 10)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)
    calls = {"large_batch": 0}

    async def split_then_succeed(*, expected_pages: int, **kwargs):
        if expected_pages >= 10:
            calls["large_batch"] += 1
            raise OcrClientError("baidu async OCR polling timeout")
        return {page_no: f"拆批后的第 {page_no} 页\n第 {page_no} 页" for page_no in range(1, expected_pages + 1)}

    service._baidu.ocr_pdf_pages_baidu_async = split_then_succeed  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")
    assert result["status"] == JobStatus.FINAL_READY.value
    assert calls["large_batch"] == 1

    diagnostics = service.store.read_diagnostics(created.job_id)
    types = [item["type"] for item in diagnostics]
    assert "baidu_split" in types


def test_run_job_auto_switches_incompatible_model_for_provider(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "compatibility.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(
        pdf_path=pdf_path,
        page_mode=PageMode.PDF,
        model_option="deepseek_ocr",
    )

    result = service.run_job(created.job_id, provider="baidu_async")
    assert result["status"] == JobStatus.FINAL_READY.value
    assert result["modelOption"] == "paddleocr_vl"

    diagnostics = service.store.read_diagnostics(created.job_id)
    warning_types = [item["type"] for item in diagnostics]
    assert "provider_warning" in warning_types


def test_low_quality_page_can_be_repaired_by_model_denoise(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "denoise.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(
        pdf_path=pdf_path,
        page_mode=PageMode.PDF,
        model_option="deepseek_ocr",
    )

    async def noisy_ocr(*args, **kwargs):
        return "delta " * 200

    async def repaired_ocr(*args, **kwargs):
        return "总序\n呈献在诸位读者面前的这套丛书。"

    service._sf.ocr_image = noisy_ocr  # type: ignore[method-assign]
    service._sf.denoise_page_text = repaired_ocr  # type: ignore[method-assign]

    result = service.run_job(created.job_id, provider="siliconflow")
    assert result["status"] == JobStatus.FINAL_READY.value

    normalized_path = service.store.page_normalized_path(created.job_id, 1)
    normalized_text = normalized_path.read_text(encoding="utf-8")
    assert "呈献在诸位读者面前" in normalized_text

    diagnostics = service.store.read_diagnostics(created.job_id)
    types = [item["type"] for item in diagnostics]
    assert "model_denoise" in types


def test_denoise_uses_configured_model_option(textferry_env: Path, make_pdf, monkeypatch) -> None:
    monkeypatch.setenv("TEXTFERRY_DENOISE_MODEL_OPTION", "paddleocr_vl")
    pdf_path = make_pdf(textferry_env / "configured-denoise-model.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(
        pdf_path=pdf_path,
        page_mode=PageMode.PDF,
        model_option="deepseek_ocr",
    )
    seen_model_keys: list[str] = []

    async def noisy_ocr(*args, **kwargs):
        return "delta " * 200

    async def repaired_ocr(*, model_option, **kwargs):
        seen_model_keys.append(model_option.key)
        return "总序\n呈献在诸位读者面前的这套丛书。"

    service._sf.ocr_image = noisy_ocr  # type: ignore[method-assign]
    service._sf.denoise_page_text = repaired_ocr  # type: ignore[method-assign]

    result = service.run_job(created.job_id, provider="siliconflow")

    assert result["status"] == JobStatus.FINAL_READY.value
    assert seen_model_keys == ["paddleocr_vl"]


def test_baidu_split_error_retries_with_rasterized_batch(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "raster-fallback.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(
        pdf_path=pdf_path,
        page_mode=PageMode.PDF,
        model_option="paddleocr_vl",
    )
    calls = {"normal": 0, "raster": 0}

    async def batch_ocr(*, pdf_path, expected_pages: int, **kwargs):
        if "raster" in str(pdf_path):
            calls["raster"] += 1
            return {page_no: f"重建 PDF 后的第 {page_no} 页\n第 {page_no} 页" for page_no in range(1, expected_pages + 1)}
        calls["normal"] += 1
        raise OcrClientError("系统错误-拆页")

    service._baidu.ocr_pdf_pages_baidu_async = batch_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")
    assert result["status"] == JobStatus.FINAL_READY.value
    assert calls["normal"] == 1
    assert calls["raster"] == 1

    diagnostics = service.store.read_diagnostics(created.job_id)
    types = [item["type"] for item in diagnostics]
    assert "baidu_raster_fallback" in types


def test_baidu_submit_connect_error_keeps_stage_and_exception_type(textferry_env: Path, make_pdf, monkeypatch) -> None:
    monkeypatch.delenv("TEXTFERRY_MOCK_OCR", raising=False)
    get_textferry_settings.cache_clear()
    get_model_registry.cache_clear()
    client = BaiduAsyncProvider()
    model = get_model_registry().options["paddleocr_vl"]
    pdf_path = make_pdf(textferry_env / "submit-error.pdf", 1)

    class RaisingAsyncClient:
        def __init__(self, *args, **kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

        async def post(self, *args, **kwargs):
            raise httpx.ConnectError("boom")

    original_client = httpx.AsyncClient
    httpx.AsyncClient = RaisingAsyncClient  # type: ignore[assignment]
    try:
        import asyncio

        async def run():
            await client.ocr_pdf_pages_baidu_async(
                pdf_path=pdf_path,
                model_option=model,
                api_key="test",
                expected_pages=1,
            )

        try:
            asyncio.run(run())
        except OcrClientError as exc:
            assert "submit" in str(exc)
            assert "ConnectError" in str(exc)
        else:
            raise AssertionError("Expected OcrClientError")
    finally:
        httpx.AsyncClient = original_client  # type: ignore[assignment]


def test_baidu_submit_non_json_response_raises_structured_error(textferry_env: Path, make_pdf, monkeypatch) -> None:
    monkeypatch.delenv("TEXTFERRY_MOCK_OCR", raising=False)
    get_textferry_settings.cache_clear()
    get_model_registry.cache_clear()
    client = BaiduAsyncProvider()
    model = get_model_registry().options["paddleocr_vl"]
    pdf_path = make_pdf(textferry_env / "submit-non-json.pdf", 1)

    class NonJsonAsyncClient:
        def __init__(self, *args, **kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

        async def post(self, *args, **kwargs):
            return httpx.Response(
                200,
                text="<html>gateway returned html</html>",
                request=httpx.Request("POST", "https://example.test/api/v2/ocr/jobs"),
            )

    original_client = httpx.AsyncClient
    httpx.AsyncClient = NonJsonAsyncClient  # type: ignore[assignment]
    try:
        import asyncio

        async def run():
            await client.ocr_pdf_pages_baidu_async(
                pdf_path=pdf_path,
                model_option=model,
                api_key="test",
                expected_pages=1,
            )

        try:
            asyncio.run(run())
        except OcrClientError as exc:
            assert exc.provider == "baidu_async"
            assert exc.stage == "submit"
            assert exc.error_type == "JSONDecodeError"
            assert "invalid JSON" in str(exc)
        else:
            raise AssertionError("Expected OcrClientError")
    finally:
        httpx.AsyncClient = original_client  # type: ignore[assignment]


def test_baidu_stage_error_writes_structured_diagnostic(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "stage-error.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(
        pdf_path=pdf_path,
        page_mode=PageMode.PDF,
        model_option="paddleocr_vl",
    )

    async def stage_error(**kwargs):
        raise OcrClientError(
            "baidu submit network error [ConnectError]: boom",
            stage="submit",
            error_type="ConnectError",
        )

    service._baidu.ocr_pdf_pages_baidu_async = stage_error  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")
    assert result["status"] == JobStatus.PARTIAL_READY.value

    diagnostics = service.store.read_diagnostics(created.job_id)
    request_errors = [item for item in diagnostics if item["type"] == "baidu_request_error"]
    assert request_errors
    assert request_errors[0]["context"]["stage"] == "submit"
    assert request_errors[0]["context"]["errorType"] == "ConnectError"


def test_baidu_submit_connect_error_aborts_without_failing_pages(textferry_env: Path, make_pdf, monkeypatch) -> None:
    monkeypatch.setenv("TEXTFERRY_BAIDU_BATCH_CONCURRENCY", "1")
    monkeypatch.setenv("TEXTFERRY_MAX_RETRIES", "1")
    pdf_path = make_pdf(textferry_env / "submit-connect-abort.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(
        pdf_path=pdf_path,
        page_mode=PageMode.PDF,
        model_option="paddleocr_vl",
    )

    async def connect_error(**kwargs):
        raise OcrClientError(
            "baidu submit network error [ConnectError]: nodename nor servname provided",
            provider="baidu_async",
            stage="submit",
            error_type="ConnectError",
            retryable=True,
        )

    service._baidu.ocr_pdf_pages_baidu_async = connect_error  # type: ignore[method-assign]

    result = service.run_job(created.job_id, provider="baidu_async")

    assert result["status"] == JobStatus.CREATED.value
    assert result["progress"]["processedPages"] == 0
    assert result["progress"]["failedPagesCount"] == 0
    assert service.store.get_page(created.job_id, 1)["status"] == PageStatus.PENDING.value
    assert service.store.get_page(created.job_id, 2)["status"] == PageStatus.PENDING.value
    event_names = [item["event"] for item in service.store.read_events(created.job_id)]
    assert "job_provider_unavailable" in event_names


def test_baidu_raster_fallback_failure_keeps_structured_stage_diagnostic(textferry_env: Path, make_pdf, monkeypatch) -> None:
    monkeypatch.setenv("TEXTFERRY_MAX_RETRIES", "1")
    pdf_path = make_pdf(textferry_env / "raster-fallback-error.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(
        pdf_path=pdf_path,
        page_mode=PageMode.PDF,
        model_option="paddleocr_vl",
    )
    calls = {"normal": 0, "raster": 0}

    async def batch_ocr(*, pdf_path, **kwargs):
        if str(pdf_path).endswith("_raster.pdf"):
            calls["raster"] += 1
            raise OcrClientError(
                "baidu result network error [ConnectError]: boom",
                stage="result",
                error_type="ConnectError",
            )
        calls["normal"] += 1
        raise OcrClientError("系统错误-拆页", stage="poll", error_type="JobFailed")

    service._baidu.ocr_pdf_pages_baidu_async = batch_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")
    assert result["status"] == JobStatus.CREATED.value
    assert all(row["status"] == "PENDING" for row in service.store.list_pages(created.job_id))
    assert calls == {"normal": 1, "raster": 1}

    diagnostics = service.store.read_diagnostics(created.job_id)
    request_errors = [item for item in diagnostics if item["type"] == "baidu_request_error"]
    assert len(request_errors) == 2
    assert request_errors[0]["context"]["stage"] == "poll"
    assert request_errors[1]["context"]["stage"] == "result"
    assert request_errors[1]["context"]["errorType"] == "ConnectError"
    assert request_errors[1]["context"]["path"] == "raster_fallback"


def test_siliconflow_request_error_keeps_stage_and_exception_type(textferry_env: Path, make_pdf, monkeypatch) -> None:
    monkeypatch.delenv("TEXTFERRY_MOCK_OCR", raising=False)
    get_textferry_settings.cache_clear()
    get_model_registry.cache_clear()
    client = SiliconFlowProvider()
    model = get_model_registry().options["deepseek_ocr"]
    image_path = textferry_env / "page.png"
    image_path.write_bytes(b"not-a-real-png")

    class RaisingAsyncClient:
        def __init__(self, *args, **kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

        async def post(self, *args, **kwargs):
            raise httpx.ConnectError("boom")

    original_client = httpx.AsyncClient
    httpx.AsyncClient = RaisingAsyncClient  # type: ignore[assignment]
    try:
        import asyncio

        async def run():
            await client.ocr_image(image_path, model, "test")

        try:
            asyncio.run(run())
        except OcrClientError as exc:
            assert "ConnectError" in str(exc)
            assert exc.stage == "request"
            assert exc.error_type == "ConnectError"
        else:
            raise AssertionError("Expected OcrClientError")
    finally:
        httpx.AsyncClient = original_client  # type: ignore[assignment]


def test_get_job_view_is_read_only(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "read-only-view.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)
    before = service.store.get_job(created.job_id)

    job_view = service.get_job_view(created.job_id)

    after = service.store.get_job(created.job_id)
    assert job_view["progress"]["totalPages"] == 2
    assert before == after


def test_job_view_exposes_structured_diagnostic_summary(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "diagnostic-summary.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(
        pdf_path=pdf_path,
        page_mode=PageMode.PDF,
        model_option="paddleocr_vl",
    )

    service.store.append_diagnostic(
        created.job_id,
        type_="baidu_request_error",
        message="baidu result network error [ConnectError]: boom",
        context={"stage": "result", "errorType": "ConnectError"},
    )
    service.store.append_diagnostic(
        created.job_id,
        type_="quality_gate",
        message="low quality OCR output",
        context={"page": 1},
    )
    service.store.append_diagnostic(
        created.job_id,
        type_="model_denoise",
        message="Denoise repaired page 1",
        context={"page": 1},
    )
    service.store.append_diagnostic(
        created.job_id,
        type_="baidu_raster_fallback",
        message="Retried batch with rasterized PDF after split-page failure",
        context={"batchNo": 1},
    )

    job_view = service.get_job_view(created.job_id)

    diagnostics = job_view["diagnostics"]
    assert diagnostics["providerStageStats"]["result"]["ConnectError"] == 1
    assert diagnostics["qualityGateStats"]["quality_gate"] == 1
    assert diagnostics["denoiseStats"]["triggered"] == 1
    assert diagnostics["batchFallbackStats"]["rasterFallbacks"] == 1


def test_final_ready_writes_machine_md_and_reader_docx(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "artifacts.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(
        pdf_path=pdf_path,
        page_mode=PageMode.BOOK,
        model_option="paddleocr_vl",
    )

    result = service.run_job(created.job_id, provider="baidu_async")

    assert result["status"] == JobStatus.FINAL_READY.value
    artifacts = result["artifacts"]
    assert artifacts["machineMd"] == "outputs/machine.md"
    assert artifacts["readerDocx"] == "outputs/reader.docx"
    assert artifacts["readerPagedDocx"] == "outputs/reader-paged.docx"
    assert artifacts["readerFlowDocx"] == "outputs/reader-flow.docx"
    assert artifacts["final"] == "outputs/reader-flow.docx"
    assert artifacts["best"] == "outputs/reader-flow.docx"

    job_dir = service.store.job_dir(created.job_id)
    machine_path = job_dir / artifacts["machineMd"]
    reader_path = job_dir / artifacts["readerDocx"]
    reader_paged_path = job_dir / artifacts["readerPagedDocx"]
    reader_flow_path = job_dir / artifacts["readerFlowDocx"]
    assert machine_path.exists()
    assert reader_path.exists()
    assert reader_paged_path.exists()
    assert reader_flow_path.exists()

    machine_text = machine_path.read_text(encoding="utf-8")
    assert "## " in machine_text
    assert "PDF页:1" in machine_text
    assert "SRC:pdf_0001" in machine_text

    reader_doc = Document(str(reader_path))
    all_text = "\n".join(paragraph.text for paragraph in reader_doc.paragraphs)
    assert "SRC:pdf_0001" not in all_text

    reader_paged_doc = Document(str(reader_paged_path))
    paged_text = "\n".join(paragraph.text for paragraph in reader_paged_doc.paragraphs)
    assert "SRC:pdf_0001" in paged_text


def test_machine_md_uses_readable_page_blocks_and_preserves_first_page_title(
    textferry_env: Path, make_pdf
) -> None:
    pdf_path = make_pdf(textferry_env / "machine-readable.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def article_ocr(*, expected_pages: int, **kwargs):
        return {
            1: (
                "复旦学报(社会科学版)\n"
                "儒家心性之学的转出\n"
                "——论子游的思想创造及其道统地位\n"
                "作者甲\n\n"
                "【摘要】这是摘要。\n"
                "【关键词】子游性自命出\n"
                "孔子以其典范性的生存开始正文。\n\n"
                "这是首页正文前半句\n"
                "后半句应该在机器版里合并。\n"
                "2022 at 1:10:10 Academic Journal Electronic Publishing House. All rights reserved http://www.cnki.net"
            ),
            2: (
                "复旦学报(社会科学版)\n"
                "儒家心性之学的转出\n"
                "第二页正文。\n"
                "Academic Journal Electronic Publishing House. All rights reserved. http://www.cnki.net"
            ),
        }

    service._baidu.ocr_pdf_pages_baidu_async = article_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    machine_text = (service.store.job_dir(created.job_id) / result["artifacts"]["machineMd"]).read_text(
        encoding="utf-8"
    )

    assert "## Page 1 | PDF页:1 | SRC:pdf_0001" in machine_text
    assert "## PDF页:1 | PDF页:1" not in machine_text
    assert "### Text" in machine_text
    assert "儒家心性之学的转出" in machine_text
    assert "儒家心性之学的转出\n\n——论子游的思想创造及其道统地位" in machine_text
    assert "【关键词】子游性自命出\n\n孔子以其典范性的生存开始正文。" in machine_text
    page_2_section = machine_text.split("## Page 2 | PDF页:2 | SRC:pdf_0002", 1)[1]
    assert "儒家心性之学的转出" not in page_2_section
    assert "这是首页正文前半句后半句应该在机器版里合并。" in machine_text
    assert "Academic Journal Electronic Publishing House" not in machine_text
    assert "http://www.cnki.net" not in machine_text


def test_reader_outputs_include_paged_and_flow_versions(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-layouts.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    result = service.run_job(created.job_id, provider="baidu_async")

    job_dir = service.store.job_dir(created.job_id)
    paged_docx = job_dir / result["artifacts"]["readerPagedDocx"]
    flow_docx = job_dir / result["artifacts"]["readerFlowDocx"]
    with ZipFile(paged_docx) as archive:
        paged_xml = archive.read("word/document.xml").decode("utf-8")
    with ZipFile(flow_docx) as archive:
        flow_xml = archive.read("word/document.xml").decode("utf-8")

    assert 'w:type="page"' in paged_xml
    assert 'w:type="page"' not in flow_xml


def test_reader_paged_uses_full_ooxml_footnote_contract(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def note_ocr(*, expected_pages: int, **kwargs):
        return {1: "正文需要说明①。\n\n① 这是脚注内容。"}

    service._baidu.ocr_pdf_pages_baidu_async = note_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    flow_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerFlowDocx"]

    with ZipFile(paged_docx) as archive:
        names = archive.namelist()
        footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")
        styles_xml = archive.read("word/styles.xml").decode("utf-8")
        settings_xml = archive.read("word/settings.xml").decode("utf-8")
        document_xml = archive.read("word/document.xml").decode("utf-8")

    assert "word/footnotes.xml" in names
    assert "这是脚注内容" in footnotes_xml
    assert "FootnoteText" in styles_xml
    assert "FootnoteReference" in styles_xml
    assert "footnotePr" in settings_xml
    assert "footnotePr" in document_xml
    assert "footnoteReference" in document_xml
    assert "superscript" in document_xml

    with ZipFile(flow_docx) as archive:
        flow_names = archive.namelist()
        flow_footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")
        flow_settings_xml = archive.read("word/settings.xml").decode("utf-8")
        flow_document_xml = archive.read("word/document.xml").decode("utf-8")

    assert "word/footnotes.xml" in flow_names
    assert "这是脚注内容" in flow_footnotes_xml
    assert 'numRestart w:val="continuous"' in flow_settings_xml
    assert 'numRestart w:val="continuous"' in flow_document_xml

    flow_doc = Document(str(flow_docx))
    flow_text = "\n".join(paragraph.text for paragraph in flow_doc.paragraphs)
    assert "① 这是脚注内容。" not in flow_text

    footnote_stats = result["diagnostics"]["footnoteStats"]
    assert footnote_stats["attemptedPages"] == 1
    assert footnote_stats["materializedNotes"] == 1
    assert footnote_stats["skippedPages"] == 0
    assert footnote_stats["variant"] == "reader-paged"


def test_reader_flow_materializes_real_footnotes_without_page_labels(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-flow-footnotes.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def note_ocr(*, expected_pages: int, **kwargs):
        return {
            1: "前句①。\n\n① 第一条脚注。",
            2: "后句②。\n\n② 第二条脚注。",
        }

    service._baidu.ocr_pdf_pages_baidu_async = note_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    flow_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerFlowDocx"]
    with ZipFile(flow_docx) as archive:
        footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")
        document_xml = archive.read("word/document.xml").decode("utf-8")
        settings_xml = archive.read("word/settings.xml").decode("utf-8")

    assert document_xml.count("footnoteReference") == 2
    assert "第一条脚注。" in footnotes_xml
    assert "第二条脚注。" in footnotes_xml
    assert 'numRestart w:val="continuous"' in settings_xml
    assert 'numRestart w:val="continuous"' in document_xml

    flow_doc = Document(str(flow_docx))
    flow_text = "\n".join(paragraph.text for paragraph in flow_doc.paragraphs)
    assert "SRC:" not in flow_text
    assert "PDF页" not in flow_text
    assert "① 第一条脚注。" not in flow_text
    assert "② 第二条脚注。" not in flow_text


def test_reader_paged_skips_non_body_page_footnotes(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes-toc.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def toc_like_ocr(*, expected_pages: int, **kwargs):
        return {1: "目录\n第一章 导论……1\n第二章 研究设计……9\n\n① 这更像目录页噪声，不应转脚注。"}

    service._baidu.ocr_pdf_pages_baidu_async = toc_like_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    with ZipFile(paged_docx) as archive:
        assert "word/footnotes.xml" not in archive.namelist()

    paged_doc = Document(str(paged_docx))
    paged_text = "\n".join(paragraph.text for paragraph in paged_doc.paragraphs)
    assert "① 这更像目录页噪声，不应转脚注。" in paged_text

    footnote_stats = result["diagnostics"]["footnoteStats"]
    assert footnote_stats["attemptedPages"] == 1
    assert footnote_stats["materializedNotes"] == 0
    assert footnote_stats["skippedPages"] == 1
    assert footnote_stats["skippedReasons"]["non_body_page"] == 1


def test_reader_paged_skips_candidates_without_inline_anchor(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes-anchor.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def no_anchor_ocr(*, expected_pages: int, **kwargs):
        return {1: "这页没有正文注释锚点。\n\n① 这是孤立的注释定义。"}

    service._baidu.ocr_pdf_pages_baidu_async = no_anchor_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    with ZipFile(paged_docx) as archive:
        assert "word/footnotes.xml" not in archive.namelist()

    footnote_stats = result["diagnostics"]["footnoteStats"]
    assert footnote_stats["attemptedPages"] == 1
    assert footnote_stats["materializedNotes"] == 0
    assert footnote_stats["skippedPages"] == 1
    assert footnote_stats["skippedReasons"]["missing_inline_anchor"] == 1


def test_reader_paged_materializes_tail_footnotes_even_without_bottom_half_margin(
    textferry_env: Path, make_pdf
) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes-tail-block.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def tail_note_ocr(*, expected_pages: int, **kwargs):
        return {
            1: (
                "成书之前①。\n"
                "后来释读又引出新的讨论②，并带来不同解释③和补充判断④。\n"
                "① 注释一。\n"
                "② 注释二。\n"
                "③ 注释三。\n"
                "④ 注释四。"
            )
        }

    service._baidu.ocr_pdf_pages_baidu_async = tail_note_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    with ZipFile(paged_docx) as archive:
        names = archive.namelist()
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")

    assert "word/footnotes.xml" in names
    assert document_xml.count("footnoteReference") == 4
    assert "注释一" in footnotes_xml
    assert "注释四" in footnotes_xml

    footnote_stats = result["diagnostics"]["footnoteStats"]
    assert footnote_stats["materializedNotes"] == 4
    assert footnote_stats["skippedPages"] == 0


def test_reader_paged_partially_materializes_when_some_note_definitions_are_unmatched(
    textferry_env: Path, make_pdf
) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes-partial-match.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def partial_note_ocr(*, expected_pages: int, **kwargs):
        return {1: "甲①。\n乙②。\n\n① 注释甲。\n② 注释乙。\n③ 注释丙。"}

    service._baidu.ocr_pdf_pages_baidu_async = partial_note_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    with ZipFile(paged_docx) as archive:
        names = archive.namelist()
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")

    assert "word/footnotes.xml" in names
    assert document_xml.count("footnoteReference") == 2
    assert "注释甲" in footnotes_xml
    assert "注释乙" in footnotes_xml

    paged_doc = Document(str(paged_docx))
    paged_text = "\n".join(paragraph.text for paragraph in paged_doc.paragraphs)
    assert "③ 注释丙。" in paged_text

    footnote_stats = result["diagnostics"]["footnoteStats"]
    assert footnote_stats["materializedNotes"] == 2
    assert footnote_stats["skippedPages"] == 0


def test_reader_paged_accepts_long_scholarly_footnotes(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes-long-note.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    long_note = (
        "李学勤：《先秦儒家著作的重大发现》；《郭店楚简与儒家经籍》，《中国哲学》第二十辑，"
        "第13、15、18页，辽宁教育出版社，1999年1月。另彭浩先生有相似的看法，认为"
        "“郭店一号墓的下葬年代约在公元前300年，距公元前316年的包山二号墓相去不久，并不会迟至战国晚期。”"
        "这条注文还继续补充不同学者的争论、版本差异、出版信息与页码来源，用来模拟扫描学术书里常见的长篇脚注。"
    )

    async def long_note_ocr(*, expected_pages: int, **kwargs):
        return {1: f"成书之前①。\n后续判断仍可参考前人研究②。\n\n① {long_note}\n② 短注。"}

    service._baidu.ocr_pdf_pages_baidu_async = long_note_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    with ZipFile(paged_docx) as archive:
        footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")
        document_xml = archive.read("word/document.xml").decode("utf-8")

    assert document_xml.count("footnoteReference") == 2
    assert "郭店楚简与儒家经籍" in footnotes_xml

    footnote_stats = result["diagnostics"]["footnoteStats"]
    assert footnote_stats["materializedNotes"] == 2
    assert footnote_stats["skippedPages"] == 0


def test_reader_paged_resets_footnote_numbering_per_page(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes-page-reset.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def page_reset_ocr(*, expected_pages: int, **kwargs):
        return {
            1: "第一页正文①。\n\n① 第一页注释。",
            2: "第二页正文①。\n\n① 第二页注释。",
        }

    service._baidu.ocr_pdf_pages_baidu_async = page_reset_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    with ZipFile(paged_docx) as archive:
        settings_xml = archive.read("word/settings.xml").decode("utf-8")
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")

    assert 'numRestart w:val="eachPage"' in settings_xml
    assert 'numRestart w:val="eachPage"' in document_xml
    assert document_xml.count("footnoteReference") == 2
    assert "第一页注释" in footnotes_xml
    assert "第二页注释" in footnotes_xml


def test_marker_normalization_recovers_fullwidth_square_bracket_note(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes-fullwidth-marker.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def fullwidth_note_ocr(*, expected_pages: int, **kwargs):
        return {1: "正文[1]。\n\n［1］ 注释内容。"}

    service._baidu.ocr_pdf_pages_baidu_async = fullwidth_note_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    with ZipFile(paged_docx) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")

    assert document_xml.count("footnoteReference") == 1
    assert "注释内容" in footnotes_xml

    footnote_stats = result["diagnostics"]["footnoteStats"]
    assert footnote_stats["materializedNotes"] == 1
    assert footnote_stats["markerRecoveredNotes"] == 1


def test_marker_normalization_recovers_circled_marker_with_ocr_noise(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes-noisy-marker.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def noisy_marker_ocr(*, expected_pages: int, **kwargs):
        return {1: "正文①。\n\n@① 注释内容。"}

    service._baidu.ocr_pdf_pages_baidu_async = noisy_marker_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    with ZipFile(paged_docx) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")

    assert document_xml.count("footnoteReference") == 1
    assert "注释内容" in footnotes_xml

    footnote_stats = result["diagnostics"]["footnoteStats"]
    assert footnote_stats["materializedNotes"] == 1
    assert footnote_stats["markerRecoveredNotes"] == 1


def test_marker_normalization_does_not_cross_match_wrong_marker(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes-wrong-marker.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def wrong_marker_ocr(*, expected_pages: int, **kwargs):
        return {1: "正文①。\n\n② 注释内容。"}

    service._baidu.ocr_pdf_pages_baidu_async = wrong_marker_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    with ZipFile(paged_docx) as archive:
        assert "word/footnotes.xml" not in archive.namelist()

    footnote_stats = result["diagnostics"]["footnoteStats"]
    assert footnote_stats["materializedNotes"] == 0
    assert footnote_stats["skippedPages"] == 1


def test_layout_assist_prefers_footer_definition_zone_when_marker_block_starts_early(
    textferry_env: Path, make_pdf
) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes-layout-assist.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def layout_assist_ocr(*, expected_pages: int, **kwargs):
        return {
            1: (
                "正文①。\n"
                "正文继续。\n"
                "① 注释一。\n"
                "② 注释二。"
            )
        }

    service._baidu.ocr_pdf_pages_baidu_async = layout_assist_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    with ZipFile(paged_docx) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")

    assert document_xml.count("footnoteReference") == 1
    assert "注释一" in footnotes_xml
    footnote_stats = result["diagnostics"]["footnoteStats"]
    assert footnote_stats["materializedNotes"] == 1
    assert footnote_stats["layoutAssistedNotes"] == 1


def test_direct_match_page_does_not_call_agent_footnote_resolution(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes-direct-no-agent.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def direct_note_ocr(*, expected_pages: int, **kwargs):
        return {1: "正文需要说明①。\n\n① 这是直接命中的脚注。"}

    async def unexpected_agent_call(*args, **kwargs):
        raise AssertionError("direct match page should not call agent footnote resolution")

    service._baidu.ocr_pdf_pages_baidu_async = direct_note_ocr  # type: ignore[method-assign]
    service._sf.resolve_footnote_conflicts = unexpected_agent_call  # type: ignore[attr-defined,method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    footnote_stats = result["diagnostics"]["footnoteStats"]
    assert footnote_stats["materializedNotes"] == 1
    assert footnote_stats["agentResolvedNotes"] == 0
    assert footnote_stats["agentRejectedNotes"] == 0


def test_agent_resolves_ambiguous_inline_anchor_when_confidence_is_high(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes-agent-high.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def ambiguous_note_ocr(*, expected_pages: int, **kwargs):
        return {1: "甲句①。\n乙句①。\n\n① 这条注释应挂到第二句。"}

    async def resolve_conflicts(*, conflicts, **kwargs):
        assert len(conflicts) == 1
        conflict = conflicts[0]
        assert conflict["conflictReason"] == "ambiguous_inline_anchor"
        assert len(conflict["anchors"]) == 2
        return [
            {
                "status": "resolved",
                "confidence": "high",
                "reason": "第二句附近的语义更匹配这条注释。",
                "matches": [{"noteMarker": "①", "anchorIndex": 1}],
            }
        ]

    service._baidu.ocr_pdf_pages_baidu_async = ambiguous_note_ocr  # type: ignore[method-assign]
    service._sf.resolve_footnote_conflicts = resolve_conflicts  # type: ignore[attr-defined,method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    flow_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerFlowDocx"]
    with ZipFile(paged_docx) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")
    with ZipFile(flow_docx) as archive:
        flow_document_xml = archive.read("word/document.xml").decode("utf-8")
        flow_footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")

    assert document_xml.count("footnoteReference") == 1
    assert "这条注释应挂到第二句" in footnotes_xml
    assert flow_document_xml.count("footnoteReference") == 1
    assert "这条注释应挂到第二句" in flow_footnotes_xml

    diagnostics = service.store.read_diagnostics(created.job_id)
    diag_types = [item["type"] for item in diagnostics]
    assert "footnote_conflict" in diag_types
    assert "footnote_agent_resolved" in diag_types

    footnote_stats = result["diagnostics"]["footnoteStats"]
    assert footnote_stats["materializedNotes"] == 1
    assert footnote_stats["agentResolvedNotes"] == 1
    assert footnote_stats["agentResolvedPages"] == 1
    assert footnote_stats["agentRejectedNotes"] == 0


def test_agent_medium_confidence_keeps_raw_footnote_text(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes-agent-medium.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def ambiguous_note_ocr(*, expected_pages: int, **kwargs):
        return {1: "甲句①。\n乙句①。\n\n① 这条注释不够稳。"}

    async def medium_conflict_result(*, conflicts, **kwargs):
        return [
            {
                "status": "resolved",
                "confidence": "medium",
                "reason": "两个位置都像，但我不够确定。",
                "matches": [{"noteMarker": "①", "anchorIndex": 0}],
            }
        ]

    service._baidu.ocr_pdf_pages_baidu_async = ambiguous_note_ocr  # type: ignore[method-assign]
    service._sf.resolve_footnote_conflicts = medium_conflict_result  # type: ignore[attr-defined,method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    with ZipFile(paged_docx) as archive:
        assert "word/footnotes.xml" not in archive.namelist()

    paged_doc = Document(str(paged_docx))
    paged_text = "\n".join(paragraph.text for paragraph in paged_doc.paragraphs)
    assert "① 这条注释不够稳。" in paged_text

    diagnostics = service.store.read_diagnostics(created.job_id)
    diag_types = [item["type"] for item in diagnostics]
    assert "footnote_conflict" in diag_types
    assert "footnote_agent_rejected" in diag_types

    footnote_stats = result["diagnostics"]["footnoteStats"]
    assert footnote_stats["materializedNotes"] == 0
    assert footnote_stats["agentResolvedNotes"] == 0
    assert footnote_stats["agentRejectedNotes"] == 1
    assert footnote_stats["agentRejectedPages"] == 1


def test_agent_parse_failure_falls_back_to_rejected_conflict(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-footnotes-agent-parse-error.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def ambiguous_note_ocr(*, expected_pages: int, **kwargs):
        return {1: "甲句①。\n乙句①。\n\n① 这条注释解析失败时也不能乱转。"}

    async def broken_agent_result(*, conflicts, **kwargs):
        raise OcrClientError(
            "footnote agent returned invalid JSON",
            provider="siliconflow",
            stage="parse",
            error_type="JSONDecodeError",
            retryable=False,
        )

    service._baidu.ocr_pdf_pages_baidu_async = ambiguous_note_ocr  # type: ignore[method-assign]
    service._sf.resolve_footnote_conflicts = broken_agent_result  # type: ignore[attr-defined,method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    with ZipFile(paged_docx) as archive:
        assert "word/footnotes.xml" not in archive.namelist()

    diagnostics = service.store.read_diagnostics(created.job_id)
    rejected_diags = [item for item in diagnostics if item["type"] == "footnote_agent_rejected"]
    assert rejected_diags
    assert rejected_diags[0]["context"]["errorType"] == "JSONDecodeError"

    footnote_stats = result["diagnostics"]["footnoteStats"]
    assert footnote_stats["materializedNotes"] == 0
    assert footnote_stats["agentRejectedNotes"] == 1


def test_flow_reader_hides_page_labels_and_merges_cross_page_sentence(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-flow-polish.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def split_sentence_ocr(*, expected_pages: int, **kwargs):
        return {
            1: "第一段因为分页被切成了前半句",
            2: "后半句应该接回同一个段落。",
        }

    service._baidu.ocr_pdf_pages_baidu_async = split_sentence_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    flow_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerFlowDocx"]
    flow_doc = Document(str(flow_docx))
    flow_text = "\n".join(paragraph.text for paragraph in flow_doc.paragraphs)

    assert "SRC:" not in flow_text
    assert "PDF页" not in flow_text
    assert "第一段因为分页被切成了前半句后半句应该接回同一个段落。" in flow_text


def test_flow_reader_removes_standalone_chinese_page_labels_but_keeps_section_markers(
    textferry_env: Path, make_pdf
) -> None:
    pdf_path = make_pdf(textferry_env / "reader-flow-cjk-page-label.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def cjk_page_label_ocr(*, expected_pages: int, **kwargs):
        return {
            1: "一三五\n凡例\n这一页正文应该保留。",
            2: "三\n这一节标题应该保留。\n这里是下一段正文。",
        }

    service._baidu.ocr_pdf_pages_baidu_async = cjk_page_label_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    flow_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerFlowDocx"]
    flow_doc = Document(str(flow_docx))
    paragraphs = [paragraph.text for paragraph in flow_doc.paragraphs if paragraph.text.strip()]
    flow_text = "\n".join(paragraphs)

    assert "一三五" not in flow_text
    assert "凡例" in flow_text
    assert "三" in paragraphs
    assert "这一节标题应该保留。" in flow_text


def test_flow_reader_only_model_polishes_suspicious_pages(textferry_env: Path, make_pdf, monkeypatch) -> None:
    monkeypatch.setenv("TEXTFERRY_ENABLE_FLOW_MODEL_POLISH", "1")
    monkeypatch.setenv("SILICONFLOW_API_KEY", "test-flow-polish-key")
    pdf_path = make_pdf(textferry_env / "reader-flow-model-polish.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def suspicious_flow_ocr(*, expected_pages: int, **kwargs):
        return {
            1: "这是正常正文第一页。",
            2: "夫民有血一三六气心知之性，而无哀乐喜怒之常。",
        }

    polish_calls: list[str] = []

    async def polish_flow_text(*, draft_text: str, model_option, api_key: str):
        polish_calls.append(draft_text)
        assert model_option.key
        assert api_key
        return draft_text.replace("一三六", "")

    service._baidu.ocr_pdf_pages_baidu_async = suspicious_flow_ocr  # type: ignore[method-assign]
    service._sf.polish_flow_text = polish_flow_text  # type: ignore[attr-defined,method-assign]

    result = service.run_job(created.job_id, provider="baidu_async")

    assert len(polish_calls) == 1
    assert "一三六" in polish_calls[0]

    flow_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerFlowDocx"]
    flow_doc = Document(str(flow_docx))
    flow_text = "\n".join(paragraph.text for paragraph in flow_doc.paragraphs if paragraph.text.strip())

    assert "一三六" not in flow_text
    assert "夫民有血气心知之性" in flow_text

    flow_polish_stats = result["diagnostics"]["flowPolishStats"]
    assert flow_polish_stats["triggered"] == 1
    assert flow_polish_stats["changed"] == 1


def test_flow_reader_does_not_merge_short_titles_into_body(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-flow-title.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def title_ocr(*, expected_pages: int, **kwargs):
        return {
            1: "李天虹著\n湖北教育出版社",
            2: "总序\n正文第一段应该另起一段。",
        }

    service._baidu.ocr_pdf_pages_baidu_async = title_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    flow_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerFlowDocx"]
    flow_doc = Document(str(flow_docx))
    paragraphs = [paragraph.text for paragraph in flow_doc.paragraphs if paragraph.text.strip()]

    assert paragraphs[:3] == ["李天虹著湖北教育出版社", "总序", "正文第一段应该另起一段。"]


def test_reader_flow_strips_unresolved_textual_criticism_note_lines(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-flow-textual-criticism.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def criticism_ocr(*, expected_pages: int, **kwargs):
        return {
            1: (
                "孟子曰：三代之得天下也以仁。\n\n"
                "①「商」，闽、监、毛本同，孔本、韩本、考文古本作「殷」。\n"
                "而巳」闽、监、毛本同，岳本、廖本、孔本、韩本、考文古本作也」。\n"
                "「为」字原无，按阮校：「不」下脱「为」字，闽、监、毛三本不脱。据补。"
            )
        }

    service._baidu.ocr_pdf_pages_baidu_async = criticism_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    paged_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerPagedDocx"]
    flow_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerFlowDocx"]

    paged_doc = Document(str(paged_docx))
    paged_text = "\n".join(paragraph.text for paragraph in paged_doc.paragraphs if paragraph.text.strip())
    assert "①「商」" in paged_text
    assert "据补" in paged_text

    flow_doc = Document(str(flow_docx))
    flow_text = "\n".join(paragraph.text for paragraph in flow_doc.paragraphs if paragraph.text.strip())
    assert "孟子曰：三代之得天下也以仁。" in flow_text
    assert "①「商」" not in flow_text
    assert "据补" not in flow_text


def test_flow_reader_separates_roman_subtitle_from_chinese_body(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "reader-flow-roman-title.pdf", 1)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    async def roman_title_ocr(*, expected_pages: int, **kwargs):
        return {1: "总序\nZONG XU\n呈献在诸位读者面前的这套丛书。"}

    service._baidu.ocr_pdf_pages_baidu_async = roman_title_ocr  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    flow_docx = service.store.job_dir(created.job_id) / result["artifacts"]["readerFlowDocx"]
    flow_doc = Document(str(flow_docx))
    paragraphs = [paragraph.text for paragraph in flow_doc.paragraphs if paragraph.text.strip()]

    assert paragraphs[:3] == ["总序", "ZONG XU", "呈献在诸位读者面前的这套丛书。"]


def test_partial_ready_machine_md_marks_failed_pages(textferry_env: Path, make_pdf) -> None:
    pdf_path = make_pdf(textferry_env / "partial-artifacts.pdf", 2)
    service = build_service()
    created = service.create_job_from_pdf_path(pdf_path=pdf_path, page_mode=PageMode.PDF)

    original_mark_done = service._mark_page_done_from_raw

    def flaky_mark_done(**kwargs):
        if kwargs["page_no"] == 2:
            raise OcrClientError("baidu result network error [ConnectError]: boom")
        return original_mark_done(**kwargs)

    service._mark_page_done_from_raw = flaky_mark_done  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    assert result["status"] == JobStatus.PARTIAL_READY.value
    artifacts = result["artifacts"]
    assert artifacts["best"] == "outputs/reader-flow.docx"
    assert artifacts["partial"] == "outputs/reader-flow.docx"
    assert artifacts["machineMd"] == "outputs/machine.md"

    machine_text = (service.store.job_dir(created.job_id) / artifacts["machineMd"]).read_text(encoding="utf-8")
    assert "PDF页:1" in machine_text
    assert "PDF页:2" in machine_text
    assert "该页识别失败" in machine_text


def test_retryable_baidu_errors_record_instability_diagnostic(textferry_env: Path, make_pdf, monkeypatch) -> None:
    monkeypatch.setenv("TEXTFERRY_BAIDU_BATCH_CONCURRENCY", "6")
    pdf_path = make_pdf(textferry_env / "adaptive-concurrency.pdf", 40)
    service = build_service()
    created = service.create_job_from_pdf_path(
        pdf_path=pdf_path,
        page_mode=PageMode.BOOK,
        model_option="paddleocr_vl",
    )
    state = {"raised": False}

    async def flaky_then_ok(*, expected_pages: int, **kwargs):
        if not state["raised"]:
            state["raised"] = True
            raise OcrClientError(
                "baidu submit network error [ConnectError]: boom",
                provider="baidu_async",
                stage="submit",
                error_type="ConnectError",
                retryable=True,
            )
        return {page_no: f"第 {page_no} 页\n正文 {page_no}" for page_no in range(1, expected_pages + 1)}

    service._baidu.ocr_pdf_pages_baidu_async = flaky_then_ok  # type: ignore[method-assign]
    result = service.run_job(created.job_id, provider="baidu_async")

    assert result["status"] == JobStatus.FINAL_READY.value
    assert result["diagnostics"]["throughputStats"]["pagesPerMinute"] > 0
    diagnostics = service.store.read_diagnostics(created.job_id)
    instability = [item for item in diagnostics if item["type"] == "baidu_concurrency_note"]
    assert instability
    assert instability[0]["context"]["unstableBatches"] >= 1


def test_retryable_baidu_errors_downgrade_later_batch_concurrency(textferry_env: Path, make_pdf, monkeypatch) -> None:
    monkeypatch.setenv("TEXTFERRY_BATCH_SIZE", "4")
    monkeypatch.setenv("TEXTFERRY_BAIDU_BATCH_SIZE", "4")
    monkeypatch.setenv("TEXTFERRY_BAIDU_BATCH_CONCURRENCY", "6")
    pdf_path = make_pdf(textferry_env / "adaptive-concurrency-downgrade.pdf", 28)
    service = build_service()
    created = service.create_job_from_pdf_path(
        pdf_path=pdf_path,
        page_mode=PageMode.PDF,
        model_option="paddleocr_vl",
    )
    pending_rows = service.store.list_pending_pages(created.job_id)
    option = service.model_registry.options["paddleocr_vl"]

    async def fake_process_batch(*, batch_no: int, **kwargs):
        return {"providerUnstable": batch_no == 1, "durationSeconds": 0.01}

    service._process_batch_baidu_async = fake_process_batch  # type: ignore[method-assign]

    import asyncio

    asyncio.run(
        service._process_pending_pages(
            created.job_id,
            pending_rows=pending_rows,
            model_option=option,
            api_key="test",
            provider="baidu_async",
            script_target=ScriptTarget.SIMP,
            progress=None,
        )
    )

    diagnostics = service.store.read_diagnostics(created.job_id)
    downgrades = [item for item in diagnostics if item["type"] == "baidu_concurrency_downgrade"]
    assert downgrades
    assert downgrades[0]["context"] == {"from": 6, "to": 3}
